"""Side-by-side review files: source and translation, paragraph by paragraph.

Built from the ``(source, translation)`` pairs collected with
:func:`core.translation.shield.record_translations`. Layout-free, so they
work for every input format and are the easiest way to check a translation.
"""

from __future__ import annotations

import html
import io
import re
from typing import Iterable, List, Optional, Tuple

Pairs = Iterable[Tuple[str, str]]

_TAG_RE = re.compile(r"<[^>]+>")
_FENCE_RE = re.compile(r"^\s*(?:```|~~~)")
_LETTER_RE = re.compile(r"[^\W\d_]")


def _display(text: str) -> str:
    """Readable text: HTML tags (OCR tables) removed, whitespace collapsed."""
    return " ".join(_TAG_RE.sub(" ", text).split())


def review_rows(pairs: Pairs) -> List[Tuple[str, str]]:
    """Readable, de-duplicated rows; units without words are dropped."""
    rows = []
    seen = set()
    for source, translation in pairs:
        if _FENCE_RE.match(source):
            continue
        row = (_display(source), _display(translation))
        if not _LETTER_RE.search(row[0]) or row in seen:
            continue
        seen.add(row)
        rows.append(row)
    return rows


_CSS = """
:root { --bg: #ffffff; --fg: #1f2328; --muted: #656d76; --line: #d0d7de;
        --head: #f6f8fa; --stripe: #fbfcfd; }
@media (prefers-color-scheme: dark) {
  :root { --bg: #0d1117; --fg: #e6edf3; --muted: #8d96a0; --line: #30363d;
          --head: #161b22; --stripe: #11161d; }
}
* { box-sizing: border-box; }
body { margin: 0; padding: 24px 16px; background: var(--bg); color: var(--fg);
       font: 15px/1.55 system-ui, -apple-system, "Segoe UI", sans-serif; }
main { max-width: 1200px; margin: 0 auto; }
h1 { font-size: 1.35rem; margin: 0 0 4px; }
p.meta { color: var(--muted); margin: 0 0 20px; }
table { width: 100%; border-collapse: collapse; table-layout: fixed; }
th, td { border-bottom: 1px solid var(--line); padding: 10px 12px;
         vertical-align: top; text-align: start; overflow-wrap: anywhere; }
th { position: sticky; top: 0; background: var(--head); font-weight: 600; }
tr:nth-child(even) td { background: var(--stripe); }
td.n, th.n { width: 3.5em; color: var(--muted); text-align: end; }
@media (max-width: 640px) {
  table, thead, tbody, tr, td { display: block; width: 100%; }
  thead { display: none; }
  td { border-bottom: none; padding: 4px 0; }
  td.n { text-align: start; padding-top: 14px; }
  tr { border-bottom: 1px solid var(--line); padding-bottom: 10px; }
}
@media print { th { position: static; } }
"""


def build_review_html(
    pairs: Pairs, *, title: str, source_language: str, target_language: str,
) -> bytes:
    rows = review_rows(pairs)
    esc = html.escape
    body = "\n".join(
        f'<tr><td class="n">{i}</td><td dir="auto">{esc(src)}</td>'
        f'<td dir="auto">{esc(tgt)}</td></tr>'
        for i, (src, tgt) in enumerate(rows, 1)
    )
    page = f"""<!doctype html>
<html><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>{esc(title)} – side by side</title><style>{_CSS}</style></head>
<body><main>
<h1>{esc(title)}</h1>
<p class="meta">{esc(source_language)} → {esc(target_language)} ·
{len(rows)} paragraphs · machine translation, please review</p>
<table><thead><tr><th class="n">#</th><th>{esc(source_language)}</th>
<th>{esc(target_language)}</th></tr></thead>
<tbody>
{body}
</tbody></table>
</main></body></html>
"""
    return page.encode("utf-8")


def build_review_docx(
    pairs: Pairs, *, title: str, source_language: str, target_language: str,
) -> Optional[bytes]:
    """A two-column Word table, or ``None`` without python-docx."""
    try:
        import docx
        from docx.shared import Pt
    except ImportError:
        return None
    rows = review_rows(pairs)
    document = docx.Document()
    document.add_heading(title, level=1)
    document.add_paragraph(
        f"{source_language} → {target_language} · machine translation, "
        "please review"
    )
    table = document.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    header = table.rows[0].cells
    header[0].text = source_language
    header[1].text = target_language
    for cell in header:
        for run in cell.paragraphs[0].runs:
            run.bold = True
    for source, translation in rows:
        cells = table.add_row().cells
        cells[0].text = source
        cells[1].text = translation
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(10)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()
