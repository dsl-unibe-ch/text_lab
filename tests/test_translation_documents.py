"""Independent PDF outputs and format-level protection integration."""

import conftest_path  # noqa: F401

from contextlib import nullcontext
import io
import json
import sys
from types import SimpleNamespace
import zipfile

import fitz
import pytest

from core import doc_ir
from core.translation import engine, gpu_profile
from core.translation import format as formats
from core.translation import pdf_workflow
from core.translation.pdf_checks import PDFIntegrityError
from core.translation.pdf_extract import extract_document
from test_translation_pdf import make_pdf


@pytest.fixture(autouse=True)
def no_gpu_services(monkeypatch):
    monkeypatch.setattr(
        engine, "translation_session", nullcontext, raising=False,
    )


def test_markdown_survives_native_pdf_failure(monkeypatch):
    def fail_pdf(*args, **kwargs):
        raise PDFIntegrityError("Layout failed.", [1])

    monkeypatch.setattr(pdf_workflow, "translate_pdf", fail_pdf)
    result = pdf_workflow.translate_pdf_outputs(
        make_pdf("native"), lambda text: text,
        stem="example.en", source_name="example.pdf", ocr_allowed=False,
    )
    assert [name for name, _ in result.outputs] == ["example.en.md"]
    assert result.blocked[0]["output"] == "Reconstructed PDF"
    assert result.blocked[0]["pages"] == [1]
    assert "Layout failed" not in result.blocked[0]["message"]
    report = json.loads(result.report_bytes())
    assert report["status"] == "partial"
    assert report["available_outputs"] == ["example.en.md"]
    assert "semantic" in report["validation_scope"]


def test_pdf_overflow_produces_both_outputs_with_a_warning():
    result = pdf_workflow.translate_pdf_outputs(
        make_pdf("native"), lambda text: "expanded words " * 300,
        stem="example.en", source_name="example.pdf", ocr_allowed=False,
    )
    assert [name for name, _ in result.outputs] == [
        "example.en.md", "example.en.pdf",
    ]
    assert not result.blocked
    assert any("small font" in warning for warning in result.warnings)


def test_native_pdf_survives_independent_markdown_failure(monkeypatch):
    def fail_markdown(*args, **kwargs):
        raise PDFIntegrityError("OCR did not finish.", [1])

    monkeypatch.setattr(
        pdf_workflow, "translate_pdf_to_markdown", fail_markdown,
    )
    result = pdf_workflow.translate_pdf_outputs(
        make_pdf("native"), lambda text: text,
        stem="example.en", source_name="example.pdf", ocr_allowed=True,
    )
    assert [name for name, _ in result.outputs] == ["example.en.pdf"]
    assert result.blocked[0]["output"] == "Markdown"


def test_mixed_pdf_blocks_both_outputs_without_ocr_capacity():
    result = pdf_workflow.translate_pdf_outputs(
        make_pdf("native", "scan"), lambda text: text,
        stem="mixed.en", source_name="mixed.pdf", ocr_allowed=False,
    )
    assert not result.outputs
    assert len(result.blocked) == 2
    assert all(issue["pages"] == [2] for issue in result.blocked)
    assert json.loads(result.report_bytes())["status"] == "blocked"


def test_native_markdown_keeps_short_and_blank_pages_without_unload(
    monkeypatch,
):
    def unexpected_unload():
        raise AssertionError("native PDFs must not evict the model")

    monkeypatch.setattr(engine, "free_translation_vram", unexpected_unload)
    markdown, assets = formats.pdf_to_markdown_bundle(
        make_pdf("native", "blank", "short"),
        free_translation_vram_first=True,
    )
    assert "complete example" in markdown and "Cover" in markdown
    assert all(f"<!-- page {number} -->" in markdown
               for number in range(1, 4))
    assert not assets
    assert formats.pdf_needs_ocr(make_pdf("short", "blank")) is False


def _stub_ocr(monkeypatch, callback):
    import core

    module = SimpleNamespace(process_document=callback)
    monkeypatch.setattr(
        gpu_profile, "sequential_ocr_allowed", lambda **kwargs: True,
    )
    monkeypatch.setitem(sys.modules, "core.auto_ocr", module)
    monkeypatch.setattr(core, "auto_ocr", module, raising=False)


def test_only_ocr_subset_is_sent_to_worker_and_original_order_is_restored(
    monkeypatch,
):
    events = []
    monkeypatch.setattr(engine, "free_translation_vram",
                        lambda: events.append("unload"))

    def recognize(path, workspace, **kwargs):
        assert kwargs["native_fast_lane"] is False
        with fitz.open(path) as subset:
            assert len(subset) == 1
        events.append("ocr")
        return doc_ir.Document(pages=[doc_ir.Page(
            page_number=1, regions=[doc_ir.Region(
                id="p1_r0", type=doc_ir.TEXT, bbox=[0, 0, 20, 20],
                reading_order=0, content={"text": "Recognized scan."},
            )],
        )])

    _stub_ocr(monkeypatch, recognize)
    document = extract_document(
        make_pdf("native", "blank", "scan", "short"),
        free_translation_vram_first=True,
    )
    assert events == ["unload", "ocr"]
    assert [page.page_number for page in document.pages] == [1, 2, 3, 4]
    assert document.pages[2].regions[0].id == "p3_r0"
    assert document.pages[2].regions[0].text == "Recognized scan."


def test_ocr_checks_free_memory_after_eviction_before_starting_worker(
    monkeypatch,
):
    events = []
    _stub_ocr(monkeypatch, lambda *args, **kwargs: events.append("worker"))
    monkeypatch.setattr(engine, "free_translation_vram",
                        lambda: events.append("evict"))

    def insufficient_memory(**kwargs):
        assert kwargs["min_free_mb"] == 10_240
        assert kwargs["device"] == "cuda:0"
        assert events == ["evict"]
        return False

    monkeypatch.setattr(
        gpu_profile, "sequential_ocr_allowed", insufficient_memory,
    )
    with pytest.raises(PDFIntegrityError, match="10 GiB") as error:
        extract_document(
            make_pdf("native", "scan"), free_translation_vram_first=True,
        )
    assert error.value.pages == (2,)
    assert events == ["evict"]


def test_omitted_ocr_page_is_not_exported_as_success(monkeypatch):
    monkeypatch.setattr(engine, "free_translation_vram", lambda: None)
    _stub_ocr(monkeypatch, lambda *args, **kwargs: doc_ir.Document(pages=[]))
    with pytest.raises(PDFIntegrityError, match="pages") as error:
        extract_document(
            make_pdf("native", "scan"), free_translation_vram_first=True,
        )
    assert error.value.pages == (2,)


def test_markdown_preserves_multiline_equations_and_page_comments():
    source = (
        "Before.\n\n$$\nA = b + c\n$$\nAfter.\n"
        "\\[\nx = y\n\\]\n<!-- page 2 -->\nFinal."
    )
    translated = formats.translate_markdown(source, str.upper)
    assert "$$\nA = b + c\n$$" in translated
    assert "\\[\nx = y\n\\]" in translated
    assert "<!-- page 2 -->" in translated
    assert "BEFORE." in translated and "FINAL." in translated


@pytest.mark.parametrize("case_sensitive,expected", [
    (True, "Berne bern"), (False, "Berne Berne"),
])
def test_glossary_case_setting_reaches_markdown_and_docx(
    case_sensitive, expected,
):
    options = {
        "glossary": {"Bern": "Berne"},
        "glossary_case_sensitive": case_sensitive,
    }
    assert formats.translate_markdown(
        "Bern bern", lambda text: text, **options,
    ) == expected
    xml = (
        '<w:document xmlns:w="http://schemas.openxmlformats.org/'
        'wordprocessingml/2006/main"><w:body><w:p><w:r>'
        '<w:t>Bern bern</w:t></w:r></w:p></w:body></w:document>'
    )
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("word/document.xml", xml)
    output = formats.translate_docx(
        buffer.getvalue(), lambda text: text, **options,
    )
    with zipfile.ZipFile(io.BytesIO(output)) as archive:
        assert expected.encode() in archive.read("word/document.xml")


def test_pdf_glossary_case_setting_is_applied():
    with fitz.open() as source:
        page = source.new_page()
        page.insert_text((40, 40), "Bern bern")
        data = source.tobytes()
    result = formats.translate_pdf(
        data, lambda text: text,
        glossary={"Bern": "Berne"}, glossary_case_sensitive=True,
    )
    with fitz.open(stream=result, filetype="pdf") as document:
        assert " ".join(document[0].get_text().split()) == "Berne bern"


def test_pptx_glossary_case_setting_is_applied(monkeypatch):
    run = SimpleNamespace(text="Bern bern")
    paragraph = SimpleNamespace(runs=[run])
    frame = SimpleNamespace(paragraphs=[paragraph])
    shape = SimpleNamespace(
        shape_type=None, has_table=False,
        has_text_frame=True, text_frame=frame,
    )
    presentation = SimpleNamespace(
        slides=[SimpleNamespace(shapes=[shape], has_notes_slide=False)],
        save=lambda buffer: buffer.write(b"validated presentation"),
    )
    monkeypatch.setitem(sys.modules, "pptx", SimpleNamespace(
        Presentation=lambda buffer: presentation,
    ))
    formats.translate_pptx(
        b"source", lambda text: text,
        glossary={"Bern": "Berne"}, glossary_case_sensitive=True,
    )
    assert run.text == "Berne bern"


def test_xlsx_glossary_case_setting_preserves_formulas():
    from openpyxl import Workbook, load_workbook

    workbook = Workbook()
    workbook.active["A1"] = "Bern bern"
    workbook.active["A2"] = "=1+1"
    buffer = io.BytesIO()
    workbook.save(buffer)
    output = formats.translate_xlsx(
        buffer.getvalue(), lambda text: text,
        glossary={"Bern": "Berne"}, glossary_case_sensitive=True,
    )
    result = load_workbook(io.BytesIO(output))
    assert result.active["A1"].value == "Berne bern"
    assert result.active["A2"].value == "=1+1"
